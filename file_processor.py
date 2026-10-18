"""
Módulo de procesamiento y extracción de texto de documentos de referencia
(PDF, DOCX, TXT) para incorporar al contexto pedagógico de la planeación.
"""

import io
from typing import Any, Dict, List
from bs4 import BeautifulSoup
import docx
try:
    import PyPDF2 as pdf_lib
except ImportError:
    import pypdf as pdf_lib
import requests

def extract_text_from_pdf(file_bytes: bytes, max_pages: int = 200) -> str:
    """Extrae texto de un archivo PDF subido en Streamlit usando PyPDF2 / pypdf."""
    try:
        reader = pdf_lib.PdfReader(io.BytesIO(file_bytes))
        total_pages = len(reader.pages)
        pages_to_read = min(total_pages, max_pages)
        
        extracted_texts = []
        for i in range(pages_to_read):
            page = reader.pages[i]
            text = page.extract_text()
            if text:
                extracted_texts.append(f"--- Página {i+1} ---\n{text.strip()}")
                
        if total_pages > max_pages:
            extracted_texts.append(f"\n[Nota: Se procesaron las primeras {max_pages} de {total_pages} páginas para optimizar la respuesta].")
            
        return "\n\n".join(extracted_texts)
    except Exception as e:
        return f"[Error al extraer contenido del PDF: {str(e)}]"

def extraer_contenidos_pda_fase5(texto: str) -> List[Dict[str, str]]:
    """
    Analiza el texto extraído del Programa Sintético de la Fase 5 (5° y 6° de primaria)
    para identificar los bloques de Campo Formativo, Contenido y PDA (5° y 6° grado).
    """
    import re
    resultados = []
    if not texto:
        return resultados

    # Expresión regular para bloques estructurados comunes en el Programa Sintético
    patron = re.compile(
        r'(?:Campo\s+Formativo[:\s]+(?P<campo>[^\n]+)\n)?'
        r'Contenido[:\s]+(?P<contenido>[^\n]+)\n'
        r'(?:PDA(?:\s+(?P<grado_pda>[56])[°o]?\s+Grado)?[:\s]+(?P<pda>[^\n]+))',
        re.IGNORECASE
    )
    for m in patron.finditer(texto):
        c_dict = m.groupdict()
        campo_raw = (c_dict.get("campo") or "Lenguajes").strip()
        campo_clean = campo_raw.lower()
        if "lenguaje" in campo_clean:
            campo_val = "Lenguajes"
        elif "saber" in campo_clean or "cient" in campo_clean:
            campo_val = "Saberes y Pensamiento Científico"
        elif "ética" in campo_clean or "etica" in campo_clean or "sociedad" in campo_clean:
            campo_val = "Ética, Naturaleza y Sociedades"
        elif "humano" in campo_clean or "comunitario" in campo_clean:
            campo_val = "De lo Humano y lo Comunitario"
        else:
            campo_val = campo_raw

        cont_val = (c_dict.get("contenido") or "").strip()
        pda_val = (c_dict.get("pda") or "").strip()
        g_val = c_dict.get("grado_pda") or "6"
        if cont_val and pda_val:
            resultados.append({
                "campo": campo_val,
                "contenido": cont_val,
                "pda": pda_val,
                "grado": f"{g_val}° de Primaria"
            })
    return resultados

def extract_text_from_docx(file_bytes: bytes, max_paragraphs: int = 150) -> str:
    """Extrae texto de un archivo Word (.docx) subido."""
    try:
        doc = docx.Document(io.BytesIO(file_bytes))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        
        # También extraer texto de tablas si existen
        table_texts = []
        for t_idx, table in enumerate(doc.tables[:10]):
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_texts.append(row_text)
                    
        combined = paragraphs[:max_paragraphs]
        if table_texts:
            combined.append("\n[Tablas en el documento]:\n" + "\n".join(table_texts[:50]))
            
        return "\n".join(combined)
    except Exception as e:
        return f"[Error al extraer contenido del archivo DOCX: {str(e)}]"

def extract_text_from_txt(file_bytes: bytes) -> str:
    """Extrae texto de un archivo de texto plano."""
    try:
        # Intentar utf-8 primero, luego latin-1
        try:
            return file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            return file_bytes.decode("latin-1", errors="replace")
    except Exception as e:
        return f"[Error al leer archivo de texto: {str(e)}]"

def process_uploaded_files(uploaded_files: List[Any]) -> Dict[str, str]:
    """Procesa una lista de archivos subidos por st.file_uploader."""
    results = {}
    if not uploaded_files:
        return results

    for uploaded_file in uploaded_files:
        filename = uploaded_file.name
        ext = filename.lower().split(".")[-1]
        file_bytes = uploaded_file.getvalue()
        
        if ext == "pdf":
            content = extract_text_from_pdf(file_bytes)
        elif ext in ["docx", "doc"]:
            content = extract_text_from_docx(file_bytes)
        elif ext in ["txt", "md", "csv"]:
            content = extract_text_from_txt(file_bytes)
        else:
            content = f"[Formato .{ext} no soportado. Por favor use .pdf, .docx o .txt]"
            
        results[filename] = content
        
    return results

def format_context_for_prompt(extracted_files: Dict[str, str], max_chars_per_file: int = 6000) -> str:
    """Da formato a los archivos extraídos para insertarlos en el prompt de Gemini."""
    if not extracted_files:
        return "No se adjuntaron archivos adicionales."
        
    formatted = []
    for filename, content in extracted_files.items():
        truncated_content = content[:max_chars_per_file]
        if len(content) > max_chars_per_file:
            truncated_content += "\n... [Fragmento condensado para el análisis pedagógico] ..."
        formatted.append(f"### ARCHIVO DE REFERENCIA: {filename}\n{truncated_content}")
        
    return "\n\n" + "\n\n".join(formatted) + "\n"


def extraer_texto_url(url: str) -> str:
    """
    Conecta al enlace ingresado usando requests y BeautifulSoup,
    extrae todos los párrafos de texto (<p>) y los devuelve como un solo bloque de texto.
    Incluye un bloque try-except por si la página bloquea la conexión o el enlace está caído,
    devolviendo un string vacío en caso de error.
    """
    if not url or not isinstance(url, str) or not url.strip():
        return ""
    try:
        headers = {
            "User-Agent": (
                "Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
                "AppleWebKit/537.36 (KHTML, like Gecko) "
                "Chrome/124.0.0.0 Safari/537.36"
            ),
            "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,*/*;q=0.8"
        }
        resp = requests.get(url.strip(), headers=headers, timeout=12)
        resp.raise_for_status()
        soup = BeautifulSoup(resp.text, "html.parser")
        parrafos = [p.get_text().strip() for p in soup.find_all("p") if p.get_text().strip()]
        return "\n\n".join(parrafos)
    except Exception:
        return ""

